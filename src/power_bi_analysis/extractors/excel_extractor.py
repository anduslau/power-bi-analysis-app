"""
Excel file extractor for .xlsx and .xls files.
Uses openpyxl for .xlsx, and maybe xlrd for .xls (future).
"""

import zipfile
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime

import openpyxl
from openpyxl.workbook import Workbook

from .base import BaseExtractor
from ..models import (
    ReportMetadata, FileType, DataSource, Column, Table,
    Parameter, Visual, Measure, Relationship
)

class ExcelExtractor(BaseExtractor):
    """Extractor for Excel files."""

    def __init__(self):
        # Map Excel data types to common SQL types
        self.type_mapping = {
            's': 'string', 'str': 'string', 'string': 'string',
            'i': 'integer', 'int': 'integer', 'whole': 'integer',
            'f': 'float', 'float': 'float', 'decimal': 'float',
            'd': 'date', 'date': 'date', 'datetime': 'datetime',
            'b': 'boolean', 'bool': 'boolean',
            'n': 'numeric', 'currency': 'numeric',
        }

    def can_extract(self, file_path: Path) -> bool:
        """Check if file is an Excel file."""
        ext = file_path.suffix.lower()
        return ext in ['.xlsx', '.xls', '.xlsm', '.xlsb']

    @property
    def supported_extensions(self) -> List[str]:
        return ['.xlsx', '.xls', '.xlsm', '.xlsb']

    def extract(self, file_path: Path) -> ReportMetadata:
        """Extract metadata from Excel file."""
        metadata = ReportMetadata(
            file_type=FileType.EXCEL,
            file_path=str(file_path),
            file_size=file_path.stat().st_size,
            extracted_at=datetime.now().isoformat()
        )

        try:
            # Load workbook
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)

            # Extract worksheets
            metadata.worksheets = [sheet.title for sheet in wb.worksheets]

            # Extract tables (Excel Tables)
            metadata.tables = self._extract_tables(wb)

            # Extract named ranges
            named_ranges = self._extract_named_ranges(wb)
            # Could convert named ranges to DataSources or add as queries
            # For now, add as queries
            for name, ref in named_ranges.items():
                metadata.queries.append({
                    'name': name,
                    'type': 'named_range',
                    'reference': ref
                })

            # Extract external connections (data sources)
            metadata.data_sources = self._extract_data_sources(wb)

            # Extract formulas? Might be heavy; skip for now.

            wb.close()

        except Exception as e:
            metadata.extraction_errors.append(f"Excel extraction error: {e}")

        return metadata

    def _extract_tables(self, wb: Workbook) -> List[Table]:
        """Extract Excel Table definitions."""
        tables = []
        for sheet in wb.worksheets:
            # In read-only mode, sheet.tables may not be available
            if not hasattr(sheet, 'tables'):
                continue
            for table in sheet.tables.values():
                # Determine columns from table range
                # For simplicity, we'll just record table name and range
                # Could parse column headers if available
                cols = []
                # Try to get column headers (first row of table)
                # Not implemented fully; placeholder
                tables.append(Table(
                    name=table.name,
                    columns=cols,
                    description=f"Excel Table in sheet '{sheet.title}', range {table.ref}",
                    source=table.ref
                ))
        return tables

    def _extract_named_ranges(self, wb: Workbook) -> Dict[str, str]:
        """Extract named ranges from workbook."""
        named_ranges = {}
        try:
            # defined_names.definedName might be a list of DefinedName objects
            # or a dict-like mapping. Handle both.
            defined_names = wb.defined_names.definedName
            if hasattr(defined_names, 'items'):
                # Dict-like
                for name, range_def in defined_names.items():
                    named_ranges[name] = str(range_def)
            else:
                # List-like
                for defined_name in defined_names:
                    named_ranges[defined_name.name] = defined_name.value
        except Exception:
            # If we can't extract named ranges, return empty dict
            pass
        return named_ranges

    def _extract_data_sources(self, wb: Workbook) -> List[DataSource]:
        """Extract external data connections."""
        data_sources = []
        # openpyxl doesn't expose connections directly.
        # Could parse workbook's connections.xml from zip.
        # For now, placeholder.
        try:
            # Attempt to read connections file from zip
            with zipfile.ZipFile(wb._archive, 'r') as z:
                if 'xl/connections.xml' in z.namelist():
                    # Parse connections XML (simplistic)
                    # Not implemented
                    pass
        except:
            pass
        return data_sources
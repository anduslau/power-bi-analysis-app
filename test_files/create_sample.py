#!/usr/bin/env python3
"""
Create a sample Excel file for testing the Excel extractor.
"""
import sys
from pathlib import Path
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

def create_sample_excel(output_path: Path):
    """Create a sample Excel file with tables, named ranges, and data."""
    wb = openpyxl.Workbook()

    # Sheet1 with a table
    ws1 = wb.active
    ws1.title = "Sales"

    # Add sample data
    data = [
        ["Region", "Product", "Sales", "Date"],
        ["North", "A", 1000, "2024-01-01"],
        ["South", "B", 1500, "2024-01-02"],
        ["East", "A", 1200, "2024-01-03"],
        ["West", "C", 1800, "2024-01-04"],
    ]
    for row in data:
        ws1.append(row)

    # Create a table
    table_range = f"A1:D{len(data)}"
    tab = Table(displayName="SalesTable", ref=table_range)
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                          showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws1.add_table(tab)

    # Add a named range
    wb.create_named_range("SalesRegion", ws1, "A2:A5")

    # Sheet2 with formulas
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Total Sales"
    ws2["B1"] = "=SUM(Sales!C2:C5)"

    # Save
    wb.save(output_path)
    print(f"Created sample Excel file at {output_path}")

if __name__ == "__main__":
    output = Path(__file__).parent / "sample.xlsx"
    create_sample_excel(output)